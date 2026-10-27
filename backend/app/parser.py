"""Document parser module providing lightweight, layered extraction.

Prioritizes fast-path extraction for digital text-based documents (pypdf, python-docx)
to completely eliminate unnecessary PyTorch/Docling initialization and prevent Windows
WinError 1450 resource exhaustion.
"""

import logging
import os
import shutil
import subprocess
import tempfile
import uuid
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Tuple

try:
    from pypdf import PdfReader  # type: ignore[import-not-found]
except ImportError:  # pragma: no cover
    PdfReader = None  # type: ignore

try:
    from docx import Document as DocxDocument  # type: ignore[import-not-found,import-untyped]
except ImportError:  # pragma: no cover
    DocxDocument = None  # type: ignore

from backend.app.config import settings
from backend.app.exceptions import DocumentParserError, FileValidationError

logger = logging.getLogger(__name__)

# Allowed extensions
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".doc"}


def clean_extracted_syllabus_text(raw_text: str) -> str:
    """
    Remove recognized repetitive syllabus boilerplate safely.
    Preserves legitimate syllabus content, topic titles, and learning items.
    """
    if not raw_text:
        return ""

    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")

    # 1. Exact recognized boilerplate phrases
    text = re.sub(
        r"(?im)^[ \t]*Unit syllabus detailed learning outcomes and lab exercises\.?[ \t]*$",
        "",
        text,
    )
    text = re.sub(
        r"(?im)^[ \t]*All rights reserved\.?[ \t]*$",
        "",
        text,
    )
    text = re.sub(
        r"(?im)^[ \t]*Internal(?: / External)? examination guidelines and grading scheme\.?[ \t]*$",
        "",
        text,
    )

    # 2. Running page numbers and margin noise
    text = re.sub(
        r"(?im)^[ \t]*Page[ \t]+\d+[ \t]+(?:of|/)[ \t]+\d+[ \t]*$",
        "",
        text,
    )
    text = re.sub(
        r"(?im)^[ \t]*Page[ \t]+\d+[ \t]*$",
        "",
        text,
    )
    text = re.sub(
        r"(?im)^[ \t]*[-–—]\s*\d+\s*[-–—][ \t]*$",
        "",
        text,
    )
    text = re.sub(
        r"(?m)^[ \t]*\d+\s*/\s*\d+[ \t]*$",
        "",
        text,
    )

    # 3. Collapse excessive blank lines
    text = re.sub(
        r"\n[ \t]*\n(?:[ \t]*\n)+",
        "\n\n",
        text,
    )

    return text.strip()


def remove_running_headers_and_footers(pages_text: List[str]) -> List[str]:
    """Detect and remove repeated running headers and footers across pages safely."""
    if len(pages_text) < 3:
        return pages_text

    from collections import Counter
    top_lines = Counter()
    bottom_lines = Counter()

    for p in pages_text:
        lines = [line.strip() for line in p.splitlines() if line.strip()]
        for l in lines[:3]:
            top_lines[l] += 1
        for l in lines[-2:]:
            bottom_lines[l] += 1

    # Candidate headers/footers appearing on >= 50% of pages and at least 3 pages
    threshold = max(3, len(pages_text) // 2)
    excluded_prefixes = ("unit", "module", "chapter", "course", "subject", "year", "sem", "topic", "#", "*", "-")
    repeated_headers = {
        line for line, count in top_lines.items()
        if count >= threshold and not any(line.lower().startswith(p) for p in excluded_prefixes)
    }
    repeated_footers = {
        line for line, count in bottom_lines.items()
        if count >= threshold and not any(line.lower().startswith(p) for p in excluded_prefixes)
    }

    if not repeated_headers and not repeated_footers:
        return pages_text

    cleaned_pages = []
    for p in pages_text:
        lines = p.splitlines()
        n = len(lines)
        filtered = []
        for idx, line in enumerate(lines):
            s = line.strip()
            if idx < 3 and s in repeated_headers:
                continue
            if idx >= max(0, n - 2) and s in repeated_footers:
                continue
            filtered.append(line)
        cleaned_pages.append("\n".join(filtered))

    return cleaned_pages


@dataclass
class ParsedDocument:
    """Normalized document representation containing extracted text and page mapping."""
    text: str
    raw_text: str = ""
    pages: List[Tuple[Optional[int], str]] = field(default_factory=list)
    has_reliable_pages: bool = False
    source_filename: str = ""
    warnings: List[str] = field(default_factory=list)

    def __post_init__(self):
        if not self.raw_text:
            self.raw_text = self.text


def validate_file_metadata(filename: str, file_size: int) -> str:
    """Validate filename extension and enforce maximum size constraint."""
    if not filename:
        raise FileValidationError("Upload filename is missing or empty.")

    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise FileValidationError(
            f"Unsupported file format '{ext}'. Allowed formats: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    if file_size <= 0:
        raise FileValidationError("Uploaded file is empty (0 bytes).")

    if file_size > settings.MAX_UPLOAD_SIZE_BYTES:
        max_mb = settings.MAX_UPLOAD_SIZE_BYTES / (1024 * 1024)
        raise FileValidationError(f"File exceeds maximum allowed size of {max_mb:.0f} MiB.")

    return ext


def parse_txt_content(content_bytes: bytes, filename: str) -> ParsedDocument:
    """Parse raw plain text with multiple encoding fallbacks."""
    text = ""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = content_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if not text.strip():
        raise DocumentParserError(f"Plain text file '{filename}' contains no readable text.")

    if len(text) > settings.MAX_EXTRACTED_TEXT_CHARS:
        logger.warning(
            "document.text_truncated filename=%s length=%d limit=%d",
            filename,
            len(text),
            settings.MAX_EXTRACTED_TEXT_CHARS,
        )
        text = text[: settings.MAX_EXTRACTED_TEXT_CHARS]

    # TXT files have no native page breaks
    cleaned_text = clean_extracted_syllabus_text(text)
    return ParsedDocument(
        text=cleaned_text,
        raw_text=text,
        pages=[(None, text)],
        has_reliable_pages=False,
        source_filename=filename,
    )


def parse_pdf_pypdf(file_path: str, filename: str) -> Optional[ParsedDocument]:
    """
    Fast, lightweight, offline PDF parsing using pypdf.
    
    Extracts text per page preserving 1-based page numbers.
    Never initializes PyTorch, layout models, or OCR.
    Returns ParsedDocument if extractable text is present, or None if the document
    is scanned or contains no machine-readable text.
    """
    if PdfReader is None:
        raise DocumentParserError("The 'pypdf' package is not installed. Please install requirements.")

    try:
        reader = PdfReader(file_path)
    except Exception as e:
        raise DocumentParserError(f"Malformed or unreadable PDF '{filename}': {str(e)}") from e

    if reader.is_encrypted:
        raise FileValidationError(f"PDF '{filename}' is encrypted or password-protected and cannot be read.")

    total_pages = len(reader.pages)
    if total_pages == 0:
        raise DocumentParserError(f"PDF '{filename}' contains zero pages.")

    if total_pages > settings.MAX_PDF_PAGES:
        raise FileValidationError(
            f"PDF exceeds maximum page limit of {settings.MAX_PDF_PAGES} pages (found {total_pages})."
        )

    pages: List[Tuple[Optional[int], str]] = []
    full_text_chunks: List[str] = []
    total_chars = 0

    for idx, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text() or ""
        except Exception as page_err:
            logger.warning("pypdf.page_extract_error page=%d err=%s", idx + 1, page_err)
            page_text = ""

        page_num = idx + 1
        pages.append((page_num, page_text))
        if page_text.strip():
            full_text_chunks.append(f"[Page {page_num}]\n{page_text}")
            total_chars += len(page_text.strip())

    full_text = "\n\n".join(full_text_chunks)

    # Determine if text is meaningful (even a short valid syllabus with just a few topics)
    if not full_text.strip() or total_chars == 0:
        logger.info(
            "document.fast_path.no_text filename=%s total_pages=%d chars=0",
            filename,
            total_pages,
        )
        return None

    if len(full_text) > settings.MAX_EXTRACTED_TEXT_CHARS:
        logger.warning(
            "document.pdf_truncated filename=%s len=%d limit=%d",
            filename,
            len(full_text),
            settings.MAX_EXTRACTED_TEXT_CHARS,
        )
        full_text = full_text[: settings.MAX_EXTRACTED_TEXT_CHARS]

    cleaned_full_text = clean_extracted_syllabus_text(full_text)
    return ParsedDocument(
        text=cleaned_full_text,
        raw_text=full_text,
        pages=pages,
        has_reliable_pages=True,
        source_filename=filename,
    )


def parse_docx_python_docx(file_path: str, filename: str) -> ParsedDocument:
    """Fast, offline DOCX parsing using python-docx extracting paragraphs and tables."""
    if DocxDocument is None:
        raise DocumentParserError("The 'python-docx' package is not installed. Please install requirements.")

    try:
        doc = DocxDocument(file_path)
    except Exception as e:
        raise DocumentParserError(f"Malformed or unreadable DOCX '{filename}': {str(e)}") from e

    chunks: List[str] = []

    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt:
            chunks.append(txt)

    for table in doc.tables:
        for row in table.rows:
            row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_cells:
                chunks.append(" | ".join(row_cells))

    full_text = "\n\n".join(chunks)
    if not full_text.strip():
        raise DocumentParserError(f"DOCX '{filename}' contains no extractable text.")

    if len(full_text) > settings.MAX_EXTRACTED_TEXT_CHARS:
        full_text = full_text[: settings.MAX_EXTRACTED_TEXT_CHARS]

    # DOCX doesn't have reliable fixed physical pages without layout engine
    cleaned_full_text = clean_extracted_syllabus_text(full_text)
    return ParsedDocument(
        text=cleaned_full_text,
        raw_text=full_text,
        pages=[(None, full_text)],
        has_reliable_pages=False,
        source_filename=filename,
    )


def parse_with_docling_if_available(file_path: str, filename: str) -> Optional[ParsedDocument]:
    """
    Attempt secondary parsing using Docling only when explicitly enabled.
    
    Guarded against Windows WinError 1450 (Insufficient system resources) and PyTorch
    allocation errors.
    """
    if not settings.ENABLE_DOCLING_FALLBACK:
        logger.info("document.docling_skipped reason=disabled_in_settings filename=%s", filename)
        return None

    logger.info("document.docling_initializing filename=%s", filename)
    try:
        from docling.document_converter import DocumentConverter  # type: ignore[import-not-found]

        converter = DocumentConverter()
        result = converter.convert(file_path)
        if hasattr(result, "document"):
            md_text = result.document.export_to_markdown()
            if md_text and md_text.strip():
                is_pdf = filename.lower().endswith(".pdf")
                cleaned_md = clean_extracted_syllabus_text(md_text)
                return ParsedDocument(
                    text=cleaned_md,
                    raw_text=md_text,
                    pages=[(1 if is_pdf else None, md_text)],
                    has_reliable_pages=is_pdf,
                    source_filename=filename,
                )
    except OSError as os_err:
        # Handles Windows [WinError 1450] Insufficient system resources
        logger.error(
            "document.docling_os_error filename=%s error=%s",
            filename,
            str(os_err),
            exc_info=True,
        )
        return None
    except Exception as e:
        logger.warning(
            "document.docling_failed filename=%s error=%s",
            filename,
            str(e),
            exc_info=True,
        )
        return None

    return None


def find_libreoffice_executable() -> Optional[str]:
    """Find soffice or libreoffice executable in system PATH or standard locations."""
    for cmd in ["soffice", "libreoffice"]:
        path = shutil.which(cmd)
        if path:
            return path
    win_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for p in win_paths:
        if os.path.exists(p):
            return p
    return None


def convert_doc_to_docx(doc_path: str, output_dir: str) -> str:
    """Convert legacy .doc file to .docx using headless LibreOffice."""
    soffice_path = find_libreoffice_executable()
    if not soffice_path:
        raise DocumentParserError(
            "LibreOffice / soffice is not installed or not found in PATH. "
            "Converting legacy .doc files requires LibreOffice. "
            "Please convert the file to .docx or install LibreOffice."
        )
    try:
        subprocess.run(
            [soffice_path, "--headless", "--convert-to", "docx", "--outdir", output_dir, doc_path],
            check=True,
            timeout=settings.DOC_CONVERSION_TIMEOUT_SECONDS,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        base_name = Path(doc_path).stem
        converted_docx = os.path.join(output_dir, f"{base_name}.docx")
        if not os.path.exists(converted_docx):
            raise DocumentParserError(f"LibreOffice conversion failed: {converted_docx} was not created.")
        return converted_docx
    except subprocess.TimeoutExpired:
        raise DocumentParserError(
            f"LibreOffice conversion timed out after {settings.DOC_CONVERSION_TIMEOUT_SECONDS} seconds."
        )
    except Exception as e:
        if isinstance(e, DocumentParserError):
            raise
        raise DocumentParserError(f"Failed to convert legacy .doc file: {str(e)}") from e


def parse_uploaded_file(
    content_bytes: bytes,
    filename: str,
) -> ParsedDocument:
    """
    Main document-parser entrypoint.
    
    1. Validates file metadata and size.
    2. Uses fast in-memory decoding for TXT.
    3. Creates uniquely named, isolated temporary files for binary formats.
    4. Fast-path: extracts digital PDFs via pypdf (zero PyTorch/Docling overhead).
    5. Fallback: only if PDF has zero extractable text and Docling is enabled, attempts OCR.
    6. Guarantees deterministic temporary directory cleanup.
    """
    logger.info("upload.accepted filename=%s size_bytes=%d", filename, len(content_bytes))
    ext = validate_file_metadata(filename, len(content_bytes))

    # Fast path for plain text: no disk I/O needed
    if ext == ".txt":
        logger.info("document.fast_path.started format=txt filename=%s", filename)
        doc = parse_txt_content(content_bytes, filename)
        logger.info("document.fast_path.completed format=txt filename=%s text_len=%d", filename, len(doc.text))
        return doc

    # Unique temporary folder and file to avoid collision across concurrent uploads
    unique_id = uuid.uuid4().hex[:10]
    temp_dir = tempfile.mkdtemp(prefix=f"syllabus_{unique_id}_")
    temp_file_path = os.path.join(temp_dir, f"doc_{unique_id}{ext}")

    try:
        with open(temp_file_path, "wb") as f:
            f.write(content_bytes)

        effective_path = temp_file_path
        effective_ext = ext

        # Legacy .doc conversion
        if ext == ".doc":
            effective_path = convert_doc_to_docx(temp_file_path, temp_dir)
            effective_ext = ".docx"

        # Fast path for digital PDF
        if effective_ext == ".pdf":
            logger.info("document.fast_path.started format=pdf filename=%s", filename)
            pdf_doc = parse_pdf_pypdf(effective_path, filename)

            if pdf_doc is not None:
                logger.info(
                    "document.fast_path.completed format=pdf filename=%s pages=%d text_len=%d",
                    filename,
                    len(pdf_doc.pages),
                    len(pdf_doc.text),
                )
                return pdf_doc

            # pypdf extracted no text: document appears to be scanned/image-only
            logger.warning(
                "document.fast_path.empty_scanned filename=%s, attempting secondary fallback",
                filename,
            )
            logger.info("document.fallback.started filename=%s", filename)
            docling_result = parse_with_docling_if_available(effective_path, filename)
            if docling_result and docling_result.text.strip():
                logger.info("document.fallback.completed filename=%s", filename)
                return docling_result

            # Both extraction paths exhausted
            raise DocumentParserError(
                f"PDF '{filename}' contains no extractable text and appears to be a scanned image. "
                "Lightweight text extraction found 0 characters, and secondary OCR processing is disabled "
                "or unavailable to prevent system resource exhaustion. Please upload a digital text-based PDF, "
                "DOCX, or TXT file."
            )

        elif effective_ext == ".docx":
            logger.info("document.fast_path.started format=docx filename=%s", filename)
            docx_doc = parse_docx_python_docx(effective_path, filename)
            logger.info("document.fast_path.completed format=docx filename=%s text_len=%d", filename, len(docx_doc.text))
            return docx_doc

        else:
            raise FileValidationError(f"Unhandled file extension: {ext}")

    finally:
        # Guarantee deterministic cleanup of temporary staging directory
        if os.path.exists(temp_dir):
            try:
                shutil.rmtree(temp_dir, ignore_errors=True)
                logger.debug("document.cleanup.completed temp_dir=%s", temp_dir)
            except Exception as cleanup_err:
                logger.warning("document.cleanup.warning temp_dir=%s err=%s", temp_dir, cleanup_err)
